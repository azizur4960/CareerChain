from io import BytesIO
from pathlib import PurePath
import re


MAX_RESUME_BYTES = 2 * 1024 * 1024
SUPPORTED_RESUME_EXTENSIONS = {".txt", ".pdf", ".docx"}


def extract_resume_text(filename: str, content: bytes) -> str:
    """Read supported resume formats in memory without storing uploaded files."""
    extension = PurePath(filename or "").suffix.lower()
    if extension not in SUPPORTED_RESUME_EXTENSIONS:
        raise ValueError("Upload a PDF, DOCX, or TXT resume.")
    if not content:
        raise ValueError("The uploaded resume is empty.")
    if len(content) > MAX_RESUME_BYTES:
        raise ValueError("Resume files must be 2 MB or smaller.")

    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")
        elif extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            from docx import Document

            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            text = "\n".join(paragraphs + table_text)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("Could not read this resume. Check the file and try again.") from error

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError("No readable text was found. Scanned resumes need OCR before upload.")
    if len(text) > 20000:
        raise ValueError("Resume text is too long. Keep it under 20,000 characters.")
    return text